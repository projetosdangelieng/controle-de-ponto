from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime
from io import BytesIO

def exportar_extrato_excel(df_extrato, funcionario_nome, data_inicio, data_fim):
    """
    Gera o extrato de um funcionário como um arquivo Excel formatado e profissional.
    Retorna (bytes_do_arquivo, nome_sugerido_do_arquivo) — pronto para um
    st.download_button, já que o servidor do Streamlit Cloud não guarda
    arquivos locais entre um clique e outro.
    """

    # Nome sugerido do arquivo
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    nome_arquivo = f"Extrato_{funcionario_nome}_{timestamp}.xlsx"

    # Criar workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Extrato"

    # === CABEÇALHO ===
    ws.merge_cells("A1:E1")
    titulo = ws["A1"]
    titulo.value = "CONTROLE DE BANCO DE HORAS"
    titulo.font = Font(size=14, bold=True, color="FFFFFF")
    titulo.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    titulo.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25

    # Informações do funcionário
    ws.merge_cells("A2:E2")
    info = ws["A2"]
    info.value = f"Funcionário: {funcionario_nome} | Período: {data_inicio} a {data_fim}"
    info.font = Font(size=11, bold=True)
    info.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Linha em branco
    ws.row_dimensions[3].height = 5

    # === CABEÇALHO DA TABELA ===
    headers = ["Data", "Tipo", "Movimentação (h)", "Saldo Acumulado (h)", "Detalhes"]
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[4].height = 20

    # === DADOS ===
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Calcular saldo acumulado
    if not df_extrato.empty:
        df_extrato = df_extrato.copy()
        df_extrato['Saldo Acumulado'] = df_extrato['Movimentacao'].cumsum()

        linha_atual = 5
        for idx, row in df_extrato.iterrows():
            ws.cell(row=linha_atual, column=1).value = row['Data']
            ws.cell(row=linha_atual, column=2).value = row['Tipo']
            ws.cell(row=linha_atual, column=3).value = row['Movimentacao']
            ws.cell(row=linha_atual, column=4).value = row['Saldo Acumulado']
            ws.cell(row=linha_atual, column=5).value = row['Detalhe'] if 'Detalhe' in row else ""

            # Aplicar formato e estilo
            for col_num in range(1, 6):
                cell = ws.cell(row=linha_atual, column=col_num)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

                # Colorir números positivos e negativos
                if col_num in [3, 4]:  # Colunas de movimentação
                    if isinstance(cell.value, (int, float)):
                        if cell.value < 0:
                            cell.font = Font(color="FF0000")  # Vermelho
                        else:
                            cell.font = Font(color="008000")  # Verde
                        cell.number_format = '#,##0.00'

            linha_atual += 1

        # === TOTALIZADORES ===
        linha_total = linha_atual + 1
        ws.merge_cells(f"A{linha_total}:B{linha_total}")
        cell_label = ws.cell(row=linha_total, column=1)
        cell_label.value = "SALDO FINAL"
        cell_label.font = Font(size=11, bold=True, color="FFFFFF")
        cell_label.fill = PatternFill(start_color="203864", end_color="203864", fill_type="solid")
        cell_label.alignment = Alignment(horizontal="right", vertical="center")
        cell_label.border = thin_border

        saldo_final = df_extrato['Saldo Acumulado'].iloc[-1]
        cell_saldo = ws.cell(row=linha_total, column=3)
        cell_saldo.value = saldo_final
        cell_saldo.font = Font(size=11, bold=True, color="FFFFFF")
        cell_saldo.fill = PatternFill(start_color="203864", end_color="203864", fill_type="solid")
        cell_saldo.alignment = Alignment(horizontal="center", vertical="center")
        cell_saldo.number_format = '#,##0.00'
        cell_saldo.border = thin_border

        ws.row_dimensions[linha_total].height = 20

    # === AJUSTAR LARGURA DAS COLUNAS ===
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 20
    ws.column_dimensions['E'].width = 30

    # Gerar em memória (nada é salvo no disco do servidor)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue(), nome_arquivo

def exportar_relatorio_consolidado_excel(dict_funcionarios):
    """
    Gera um relatório consolidado de todos os funcionários como um arquivo Excel.
    dict_funcionarios: dicionário {nome_funcionario: dataframe_extrato}
    Retorna (bytes_do_arquivo, nome_sugerido_do_arquivo).
    """

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    nome_arquivo = f"Relatorio_Consolidado_{timestamp}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Consolidado"

    # === CABEÇALHO ===
    ws.merge_cells("A1:F1")
    titulo = ws["A1"]
    titulo.value = "RELATÓRIO CONSOLIDADO - BANCO DE HORAS"
    titulo.font = Font(size=14, bold=True, color="FFFFFF")
    titulo.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    titulo.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25

    info = ws["A2"]
    info.value = f"Gerado em: {datetime.now().strftime('%d/%m/%Y às %H:%M:%S')}"
    info.font = Font(size=10, italic=True)

    # === CABEÇALHO DA TABELA ===
    headers = ["Funcionário", "Saldo Total (h)", "Registros", "Ajustes", "Últimas Movimentações"]
    linha = 4
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=linha, column=col_num)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[linha].height = 20

    # === DADOS ===
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    linha_dados = 5
    for funcionario, df in dict_funcionarios.items():
        if not df.empty:
            saldo_total = df['Movimentacao'].sum()
            total_registros = len(df[df['Tipo'] == 'Registro Diário'])
            total_ajustes = len(df[df['Tipo'] == 'Ajuste/Saque'])
            ultimas_mov = df['Detalhe'].iloc[-1] if 'Detalhe' in df.columns else "N/A"

            ws.cell(row=linha_dados, column=1).value = funcionario
            ws.cell(row=linha_dados, column=2).value = saldo_total
            ws.cell(row=linha_dados, column=3).value = total_registros
            ws.cell(row=linha_dados, column=4).value = total_ajustes
            ws.cell(row=linha_dados, column=5).value = str(ultimas_mov)[:50]  # Primeiros 50 caracteres

            for col_num in range(1, 6):
                cell = ws.cell(row=linha_dados, column=col_num)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

                if col_num == 2:  # Coluna de saldo
                    if saldo_total < 0:
                        cell.font = Font(color="FF0000")
                    else:
                        cell.font = Font(color="008000")
                    cell.number_format = '#,##0.00'

            linha_dados += 1

    # === AJUSTAR LARGURA DAS COLUNAS ===
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 40

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue(), nome_arquivo
